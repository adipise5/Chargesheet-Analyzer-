"""Extract text from Word (.docx and .doc) documents."""
from __future__ import annotations

import re
from pathlib import Path


def extract_docx_text(path: Path) -> str:
    """Extract all paragraph text from a .docx file."""
    from docx import Document
    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))
    return "\n\n".join(paragraphs)


def extract_doc_text(path: Path) -> str:
    """Extract text from a legacy .doc (OLE) file using olefile."""
    import olefile
    if not olefile.isOleFile(str(path)):
        raise ValueError("Not a valid .doc (OLE) file")
    ole = olefile.OleFileIO(str(path))
    try:
        if ole.exists("WordDocument"):
            stream = ole.openstream("WordDocument")
            raw = stream.read()
            # Extract ASCII/UTF-8 text runs from the binary stream
            # This is a best-effort extraction; complex formatting may be lost
            text = _extract_text_from_word_binary(raw)
            if text.strip():
                return text
        # Fallback: try to read any text stream
        for entry in ole.listdir():
            name = "/".join(entry)
            if "text" in name.lower() or "word" in name.lower():
                try:
                    data = ole.openstream(entry).read()
                    decoded = data.decode("utf-8", errors="ignore")
                    if len(decoded.strip()) > 20:
                        return decoded
                except Exception:
                    continue
    finally:
        ole.close()
    raise ValueError("Could not extract text from .doc file")


def _extract_text_from_word_binary(raw: bytes) -> str:
    """Best-effort text extraction from Word binary format."""
    # Try UTF-16LE decoding (common in .doc files)
    try:
        text_utf16 = raw.decode("utf-16-le", errors="ignore")
        # Filter to printable characters and common punctuation
        cleaned = re.sub(r"[^\x20-\x7E\n\r\t\u0A80-\u0AFF.,;:!?()'\"-/]", " ", text_utf16)
        cleaned = re.sub(r" {3,}", "\n", cleaned)
        cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
        if len(cleaned.strip()) > 50:
            return cleaned.strip()
    except Exception:
        pass
    # Fallback to ASCII extraction
    text_ascii = raw.decode("ascii", errors="ignore")
    cleaned = re.sub(r"[^\x20-\x7E\n\r\t]", " ", text_ascii)
    cleaned = re.sub(r" {3,}", "\n", cleaned)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()


def extract_word_text(path: Path) -> str:
    """Extract text from a Word document (.doc or .docx)."""
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return extract_docx_text(path)
    elif suffix == ".doc":
        return extract_doc_text(path)
    else:
        raise ValueError(f"Unsupported Word file format: {suffix}")
